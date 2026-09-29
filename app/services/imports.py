"""Admin lead imports with caller-owned transactions and notification delivery."""
import csv
import io
from dataclasses import dataclass
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.exc import DataError, IntegrityError

from app.constants import PHONE_LABELS
from app.models import User
from app.schemas.leads import LeadCreateSchema
from app.services.base import TenantService
from app.services.leads import LeadService
from app.utils.phone_utils import clean_phone_number

FIELDS = frozenset(('name', 'contact_person', 'contact_title', 'email', 'phone',
    'phone_label', 'secondary_phone', 'secondary_phone_label', 'address', 'city',
    'state', 'zip', 'notes', 'type', 'lead_status', 'lead_source'))
TEMPLATE = ('Company Name,Contact Person,Contact Title,Email,Phone,Phone Label,'
            'Secondary Phone,Secondary Phone Label,Address,City,State,Zip,Notes,Type,Lead Status\n')
MAX_COLUMNS = 200
MAX_EXPANDED_BYTES = 100 * 1024 * 1024


@dataclass(frozen=True)
class ImportResult:
    response: dict
    notification: dict | None


class ImportService(TenantService):
    def __init__(self, session, principal, *, max_rows=10000, max_bytes=20*1024*1024):
        super().__init__(session, principal)
        self._require_admin()
        self.max_rows = max_rows
        self.max_bytes = max_bytes

    def template(self):
        return TEMPLATE

    def _read(self, upload):
        if upload is None:
            raise ValueError('No file uploaded')
        filename = (upload.filename or '').lower()
        if not filename.endswith(('.csv', '.xlsx')):
            raise ValueError('Unsupported file format')
        upload.stream.seek(0)
        data = upload.stream.read(self.max_bytes + 1)
        if len(data) > self.max_bytes:
            raise ValueError('Import file exceeds size limit')
        workbook = None
        try:
            if filename.endswith('.csv'):
                for encoding in ('utf-8-sig', 'cp1252', 'latin1'):
                    try:
                        decoded = data.decode(encoding)
                        break
                    except UnicodeDecodeError:
                        continue
                rows = csv.reader(io.StringIO(decoded, newline=''), strict=True)
            else:
                with ZipFile(io.BytesIO(data)) as archive:
                    members = archive.infolist()
                    if len(members) > 1000 or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                        raise ValueError('Expanded spreadsheet exceeds size limit')
                workbook = load_workbook(io.BytesIO(data), read_only=True,
                                         data_only=True, keep_links=False)
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise ValueError('Import has too many columns')
                rows = sheet.iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else '' for v in next(rows, [])]
            if not headers or len(headers) > MAX_COLUMNS or any(not h for h in headers):
                raise ValueError('Import requires nonempty headers and at most 200 columns')
            if len(set(headers)) != len(headers):
                raise ValueError('Import column headers must be unique')
            values = []
            for row in rows:
                if not row:  # Match CSV blank-line behavior.
                    continue
                if len(row) > len(headers):
                    raise ValueError('Import row has more values than headers')
                values.append([str(v) if v is not None else '' for v in row]
                              + [''] * (len(headers) - len(row)))
                if len(values) > self.max_rows:
                    raise ValueError('Import exceeds row limit')
            return headers, values
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Could not read import file') from exc
        finally:
            if workbook is not None:
                workbook.close()

    def preview(self, upload):
        headers, rows = self._read(upload)
        return {'headers': headers, 'rows': rows[:10], 'totalRows': len(rows)}

    @staticmethod
    def _mappings(mappings, headers):
        if not isinstance(mappings, list) or len(mappings) > MAX_COLUMNS:
            raise ValueError('Invalid column mappings')
        selected, targets, sources = [], set(), set()
        for mapping in mappings:
            if not isinstance(mapping, dict):
                raise ValueError('Invalid column mapping')
            source, target = mapping.get('csvColumn'), mapping.get('leadField')
            if not isinstance(source, str) or not isinstance(target, str):
                raise ValueError('Invalid column mapping')
            if not target:
                continue
            if target not in FIELDS or source not in headers:
                raise ValueError('Unknown import column or lead field')
            if target in targets or source in sources:
                raise ValueError('Duplicate column mapping')
            targets.add(target)
            sources.add(source)
            selected.append((headers.index(source), target))
        if 'name' not in targets:
            raise ValueError("'name' field (Company Name) is required")
        return selected

    def submit(self, upload, mappings, assigned_email):
        if not isinstance(assigned_email, str) or not assigned_email.strip():
            raise ValueError('Assigned user not found or inactive')
        assignee = self._query(User).filter(User.email == assigned_email.strip(),
                                           User.is_active.is_(True)).first()
        if assignee is None:
            raise ValueError('Assigned user not found or inactive')
        actor = self._query(User).filter(User.id == self.principal.user_id,
                                        User.is_active.is_(True)).first()
        if actor is None:
            raise ValueError('Importing user not found or inactive')
        headers, rows = self._read(upload)
        selected = self._mappings(mappings, headers)
        # sqlite3 legacy transaction mode does not BEGIN for reads/SAVEPOINT.
        # Ensure released row savepoints remain under the caller's rollback.
        connection = self.session.connection()
        if (connection.dialect.name == 'sqlite'
                and not connection.connection.driver_connection.in_transaction):
            connection.exec_driver_sql('BEGIN')
        lead_service = LeadService(self.session, self.principal)
        names, failures, warnings = [], [], []
        for row_number, row in enumerate(rows, start=2):
            try:
                fields = {}
                for position, target in selected:
                    value = row[position].strip()
                    if not value:
                        continue
                    if target in ('phone', 'secondary_phone'):
                        value = clean_phone_number(value)
                        if not value:
                            warnings.append(f'Invalid phone on row {row_number}')
                            continue
                    elif target == 'email':
                        value = value.lower()
                    elif target.endswith('_label'):
                        canonical = next((p for p in PHONE_LABELS if p.lower() == value.lower()), None)
                        if canonical is None:
                            warnings.append(f"Unknown phone label '{value}' on row {row_number}")
                        value = canonical or 'work'
                    fields[target] = value
                if not fields.get('name'):
                    raise ValueError("Missing required 'name' field")
                if 'secondary_phone' in fields and 'secondary_phone_label' not in fields:
                    fields['secondary_phone_label'] = 'mobile'
                validated = LeadCreateSchema(**fields)
                with self.session.begin_nested():
                    lead_service.create(validated, assigned_to=assignee.id)
                names.append(validated.name)
            except ValidationError as exc:
                message = '; '.join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}"
                                    for e in exc.errors(include_input=False, include_url=False,
                                                        include_context=False))
                failures.append({'row': row_number, 'data': dict(zip(headers, row)), 'error': message})
            except (ValueError, IntegrityError, DataError) as exc:
                message = 'Row violates a data constraint' if isinstance(exc, (IntegrityError, DataError)) else str(exc)
                failures.append({'row': row_number, 'data': dict(zip(headers, row)), 'error': message})
            # Operational/programming failures propagate; adapters must roll back
            # the entire batch instead of reporting a misleading partial success.
        count, failed = len(names), len(failures)
        notification = None
        if count:
            summary = '\n'.join(f'- {name}' for name in names[:10])
            more = f'\n...and {count - 10} more.' if count > 10 else ''
            notification = {'subject': 'New Leads Assigned to You', 'recipient': assignee.email,
                'body': f"You've been assigned {count} new leads from a recent import by {actor.email}.\n\n"
                        f'Sample of assigned leads:\n{summary}{more}\n\n'
                        'Please log in to the CRM to view all your leads.'}
        return ImportResult({'message': f'Import complete: {count} succeeded, {failed} failed.',
            'successful_imports': count, 'failed_imports': failed,
            'warnings': list(dict.fromkeys(warnings)), 'failures': failures}, notification)
